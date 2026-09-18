"""Rebuild the model-optimization section in the identification note.

The original document before the previous optimization-section insertions is
kept alongside the report as a recovery copy.  This script uses that unchanged
copy as the base so the requested three-part organization is clean and the
original report sections remain untouched.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOCX_PATH = PROJECT_ROOT / "新峰测验" / "260908 识别模型实现说明.docx"
SOURCE_PATH = DOCX_PATH.with_name(DOCX_PATH.name + ".round2-backup")
ASSET_DIR = PROJECT_ROOT / "IWaste_rec_model" / "output" / "round2_docx_figures"
FONT_PATH = Path(r"C:\Windows\Fonts\msyh.ttc")

COLORS = {
    "blue": "#4E79A7",
    "green": "#59A14F",
    "orange": "#F28E2B",
    "purple": "#B07AA1",
    "red": "#E15759",
    "grid": "#D9E2F3",
    "text": "#1F2933",
    "axis": "#52616B",
}


def cn_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_PATH), size=size, index=1 if bold else 0)


def rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def centered(draw: ImageDraw.ImageDraw, box: tuple[float, float, float, float], text: str, ft, color: str) -> None:
    left, top, right, bottom = box
    bb = draw.textbbox((0, 0), text, font=ft)
    draw.text(
        (left + (right - left - (bb[2] - bb[0])) / 2, top + (bottom - top - (bb[3] - bb[1])) / 2),
        text,
        font=ft,
        fill=rgb(color),
    )


def simple_bar_chart(
    path: Path,
    title: str,
    labels: list[str],
    values: list[float],
    colors: list[str],
    footer: str,
    ymin: float = 0,
) -> None:
    image = Image.new("RGB", (1800, 690), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), title, font=cn_font(34, True), fill=rgb(COLORS["text"]))
    left, right, top, bottom = 140, 1730, 125, 560
    yspan = 100 - ymin
    ticks = list(range(int(ymin), 101, 5 if ymin >= 70 else 20))
    for tick in ticks:
        y = bottom - (bottom - top) * (tick - ymin) / yspan
        draw.line((left, y, right, y), fill=rgb(COLORS["grid"]), width=2)
        draw.text((62, y - 13), f"{tick}", font=cn_font(18), fill=rgb(COLORS["axis"]))
    draw.line((left, top, left, bottom), fill=rgb(COLORS["axis"]), width=2)
    draw.line((left, bottom, right, bottom), fill=rgb(COLORS["axis"]), width=2)
    step = (right - left) / len(values)
    width = min(165, step * 0.52)
    for i, (label, value, color) in enumerate(zip(labels, values, colors, strict=True)):
        cx = left + step * (i + 0.5)
        y = bottom - (bottom - top) * (value - ymin) / yspan
        draw.rounded_rectangle((cx - width / 2, y, cx + width / 2, bottom), radius=7, fill=rgb(color))
        centered(draw, (cx - 105, y - 43, cx + 105, y - 3), f"{value:.2f}%", cn_font(19, True), COLORS["text"])
        centered(draw, (cx - step / 2, bottom + 15, cx + step / 2, bottom + 58), label, cn_font(19), COLORS["text"])
    centered(draw, (0, 615, 1800, 665), footer, cn_font(19), COLORS["axis"])
    image.save(path, quality=95)


def grouped_material_chart(path: Path, title: str, series: list[tuple[str, list[float], str]], footer: str) -> None:
    image = Image.new("RGB", (1800, 760), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), title, font=cn_font(34, True), fill=rgb(COLORS["text"]))
    left, right, top, bottom = 145, 1730, 125, 605
    ymin = 70
    for tick in range(70, 101, 5):
        y = bottom - (bottom - top) * (tick - ymin) / (100 - ymin)
        draw.line((left, y, right, y), fill=rgb(COLORS["grid"]), width=2)
        draw.text((65, y - 13), str(tick), font=cn_font(18), fill=rgb(COLORS["axis"]))
    draw.line((left, top, left, bottom), fill=rgb(COLORS["axis"]), width=2)
    draw.line((left, bottom, right, bottom), fill=rgb(COLORS["axis"]), width=2)
    materials = ["煤矸石", "矿渣", "钢渣"]
    group_step = (right - left) / 3
    bar_width = min(100, group_step / (len(series) + 1))
    for group, material in enumerate(materials):
        center = left + group_step * (group + 0.5)
        offset0 = -(len(series) - 1) * bar_width / 2
        for j, (_, values, color) in enumerate(series):
            cx = center + offset0 + j * bar_width
            value = values[group]
            y = bottom - (bottom - top) * (value - ymin) / (100 - ymin)
            draw.rounded_rectangle((cx - bar_width * 0.36, y, cx + bar_width * 0.36, bottom), radius=5, fill=rgb(color))
            centered(draw, (cx - 56, y - 35, cx + 56, y - 2), f"{value:.1f}", cn_font(16, True), COLORS["text"])
        centered(draw, (center - 130, bottom + 16, center + 130, bottom + 58), material, cn_font(22, True), COLORS["text"])
    legend_x = 440
    for label, _, color in series:
        draw.rectangle((legend_x, 674, legend_x + 28, 700), fill=rgb(color))
        draw.text((legend_x + 38, 671), label, font=cn_font(18), fill=rgb(COLORS["text"]))
        legend_x += 250
    centered(draw, (0, 710, 1800, 750), footer, cn_font(17), COLORS["axis"])
    image.save(path, quality=95)


def confusion_chart(path: Path) -> None:
    image = Image.new("RGB", (1800, 690), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 35), "裁剪与视角实验：总体准确率及煤矸石误识为钢渣数量", font=cn_font(34, True), fill=rgb(COLORS["text"]))
    left1, right1, left2, right2, top, bottom = 120, 900, 1030, 1710, 135, 555
    # accuracy panel
    for tick in range(85, 101, 5):
        y = bottom - (bottom - top) * (tick - 85) / 15
        draw.line((left1, y, right1, y), fill=rgb(COLORS["grid"]), width=2)
        draw.text((48, y - 13), str(tick), font=cn_font(18), fill=rgb(COLORS["axis"]))
    draw.line((left1, top, left1, bottom), fill=rgb(COLORS["axis"]), width=2)
    draw.line((left1, bottom, right1, bottom), fill=rgb(COLORS["axis"]), width=2)
    labels = ["V0\n六视图 0.78", "V1\n全图", "V2\n全图+中心", "V3\n六视图 0.90"]
    acc = [96.97, 91.77, 94.37, 94.81]
    err = [5, 12, 9, 6]
    cols = [COLORS["green"], COLORS["blue"], COLORS["orange"], COLORS["purple"]]
    step = (right1 - left1) / 4
    for i, value in enumerate(acc):
        cx = left1 + step * (i + 0.5)
        y = bottom - (bottom - top) * (value - 85) / 15
        draw.rounded_rectangle((cx - 58, y, cx + 58, bottom), radius=6, fill=rgb(cols[i]))
        centered(draw, (cx - 82, y - 40, cx + 82, y - 3), f"{value:.2f}%", cn_font(17, True), COLORS["text"])
        centered(draw, (cx - step / 2, bottom + 12, cx + step / 2, bottom + 58), labels[i], cn_font(15), COLORS["text"])
    draw.text((270, 585), "（a）独立测试集原始准确率", font=cn_font(21, True), fill=rgb(COLORS["text"]))
    # confusion panel
    for tick in range(0, 13, 3):
        y = bottom - (bottom - top) * tick / 12
        draw.line((left2, y, right2, y), fill=rgb(COLORS["grid"]), width=2)
        draw.text((970, y - 13), str(tick), font=cn_font(18), fill=rgb(COLORS["axis"]))
    draw.line((left2, top, left2, bottom), fill=rgb(COLORS["axis"]), width=2)
    draw.line((left2, bottom, right2, bottom), fill=rgb(COLORS["axis"]), width=2)
    step = (right2 - left2) / 4
    for i, value in enumerate(err):
        cx = left2 + step * (i + 0.5)
        y = bottom - (bottom - top) * value / 12
        draw.rounded_rectangle((cx - 58, y, cx + 58, bottom), radius=6, fill=rgb(cols[i]))
        centered(draw, (cx - 60, y - 40, cx + 60, y - 3), str(value), cn_font(19, True), COLORS["text"])
        centered(draw, (cx - step / 2, bottom + 12, cx + step / 2, bottom + 58), ["V0", "V1", "V2", "V3"][i], cn_font(17), COLORS["text"])
    draw.text((1160, 585), "（b）煤矸石 → 钢渣错误张数", font=cn_font(21, True), fill=rgb(COLORS["text"]))
    image.save(path, quality=95)


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    tcpr.append(element)


def set_run(run, size: float = 10.5, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_before(document: Document, anchor, text: str = "", style: str | None = None, align=None):
    paragraph = document.add_paragraph(style=style)
    if text:
        set_run(paragraph.add_run(text))
    if align is not None:
        paragraph.alignment = align
    anchor._p.addprevious(paragraph._p)
    return paragraph


def add_caption(document: Document, anchor, text: str) -> None:
    paragraph = add_before(document, anchor, text, align=WD_ALIGN_PARAGRAPH.CENTER)
    for run in paragraph.runs:
        set_run(run, size=10)


def add_figure(document: Document, anchor, figure: Path, caption: str) -> None:
    paragraph = add_before(document, anchor, align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph.add_run().add_picture(str(figure), width=Cm(15.2))
    add_caption(document, anchor, caption)


def add_table(document: Document, anchor, headers: list[str], rows: list[list[str]], caption: str, widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers, strict=True)):
        cell.text = ""
        set_run(cell.paragraphs[0].add_run(text), size=8.8, bold=True, color="FFFFFF")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade(cell, "1F4E78")
        if widths:
            cell.width = Cm(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, text) in enumerate(zip(cells, values, strict=True)):
            cell.text = ""
            set_run(cell.paragraphs[0].add_run(text), size=8.6, bold=(row_index == len(rows) - 1 and "平均" in values[0]))
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                shade(cell, "EAF2F8")
            if widths:
                cell.width = Cm(widths[i])
    anchor._p.addprevious(table._tbl)
    add_caption(document, anchor, caption)


def main() -> None:
    if not SOURCE_PATH.exists():
        raise FileNotFoundError(f"missing source backup: {SOURCE_PATH}")
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    aug_figure = ASSET_DIR / "optimization_augmentation.png"
    proto_figure = ASSET_DIR / "optimization_prototype.png"
    view_figure = ASSET_DIR / "optimization_view_crop.png"
    simple_bar_chart(
        aug_figure,
        "图像增强实验：固定独立测试集上的总体原始识别准确率",
        ["基线", "增强\nseed 42", "增强\nseed 2026", "增强\nseed 2027", "增强\n三次平均"],
        [92.64, 96.97, 94.81, 94.37, 95.38],
        [COLORS["blue"], COLORS["green"], COLORS["orange"], COLORS["purple"], COLORS["red"]],
        "训练集：基线 990 张；增强组 1,980 张。验证集 143 张、测试集 231 张均保持不变。",
        ymin=85,
    )
    grouped_material_chart(
        proto_figure,
        "原型数量实验：不同物料的原始识别准确率",
        [
            ("每类 1 个原型", [83.10, 80.68, 100.00], COLORS["blue"]),
            ("每类 3 个原型", [92.96, 98.86, 98.61], COLORS["green"]),
            ("每类 5 个原型", [94.37, 80.68, 98.61], COLORS["orange"]),
        ],
        "固定使用增强 seed 42 数据集、六视图 0.78 裁剪策略；柱顶为准确率（%）。",
    )
    confusion_chart(view_figure)

    document = Document(SOURCE_PATH)
    conclusion = next((p for p in document.paragraphs if p.text.strip() == "结论"), None)
    if conclusion is None:
        raise ValueError("cannot locate the conclusion heading")

    heading = add_before(document, conclusion, "两阶段采集图片模型优化实验", style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_before(
        document,
        conclusion,
        "在原有第二阶段现场识别测试的基础上，进一步将第一阶段和第二阶段料棚图片按采集批次重建为三分类数据集。除图像增强的随机种子外，各组实验均固定独立验证集和测试集，避免同一连续采集片段同时进入训练与测试；原始识别准确率用于比较模型直接分类能力，拒识率和正确接受率仅用于说明阈值校准后的输出变化。以下实验依次考察训练图像增强、每类原型数量以及裁剪和视角设置对模型的影响。",
    )

    h2 = add_before(document, conclusion, "图像增强实验", style="Heading 2")
    h2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_before(
        document,
        conclusion,
        "基线模型采用 DINOv2 特征提取、每类 3 个原型、六视图投票和 0.78 裁剪比例，训练、验证和测试图片分别为 990、143 和 231 张。增强组保持类别、批次划分、模型结构和测试集不变，仅对训练原图施加轻度随机裁剪、水平翻转、小角度旋转、亮度与对比度调整及轻度模糊，每张训练原图生成 1 张增强图。为降低单次随机增强带来的偶然性，采用 42、2026、2027 三个随机种子重复实验。",
    )
    add_table(
        document,
        conclusion,
        ["实验组", "训练图片数", "原始准确率", "拒识率", "正确接受率"],
        [
            ["基线模型", "990", "92.64%", "26.41%", "71.86%"],
            ["增强 seed 42", "1,980", "96.97%", "21.21%", "76.62%"],
            ["增强 seed 2026", "1,980", "94.81%", "22.08%", "75.76%"],
            ["增强 seed 2027", "1,980", "94.37%", "22.94%", "74.89%"],
            ["增强组平均值", "1,980", "95.38%", "22.08%", "75.76%"],
        ],
        "表4 图像增强实验的独立测试结果（验证集 n=143，测试集 n=231）",
    )
    add_figure(document, conclusion, aug_figure, "图9 图像增强实验在固定独立测试集上的总体原始识别准确率")
    add_before(
        document,
        conclusion,
        "三次增强实验的原始准确率均高于基线，平均值由 92.64% 提升至 95.38%，提高 2.74 个百分点；拒识率平均降低 4.33 个百分点，正确接受率提高 3.90 个百分点。按物料看，煤矸石、矿渣和钢渣的增强实验平均准确率分别为 92.02%、97.35% 和 96.29%，相对基线分别提高 1.88、1.90 和 4.62 个百分点。增强后的主要残余错误仍为煤矸石被识别为钢渣（每次 5 至 6 张），说明轻度增强可提高整体稳定性，但尚不能消除两类物料的外观相似问题。由于三次结果在 94.37% 至 96.97% 之间波动，后续报告以三次平均值而非单次最高值作为增强结论。",
    )

    h2 = add_before(document, conclusion, "原型数量实验", style="Heading 2")
    h2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_before(
        document,
        conclusion,
        "在图像增强 seed 42 数据集上，固定 DINOv2 特征、六视图 0.78 裁剪和其他推理参数，仅将每类原型数量设置为 1、3 和 5。原型可理解为同一物料的若干典型外观代表：数量过少可能不能覆盖现场差异，数量过多则可能把局部偶然外观也当作典型特征。",
    )
    add_table(
        document,
        conclusion,
        ["每类原型数", "原始准确率", "拒识率", "正确接受率", "煤矸石", "矿渣", "钢渣"],
        [
            ["1", "87.45%", "38.53%", "57.58%", "83.10%", "80.68%", "100.00%"],
            ["3", "96.97%", "21.21%", "76.62%", "92.96%", "98.86%", "98.61%"],
            ["5", "90.48%", "18.61%", "76.19%", "94.37%", "80.68%", "98.61%"],
        ],
        "表5 原型数量实验的独立测试结果（后三列为各物料原始识别准确率）",
    )
    add_figure(document, conclusion, proto_figure, "图10 不同原型数量下三种物料的原始识别准确率")
    add_before(
        document,
        conclusion,
        "每类 3 个原型时总体原始准确率最高，为 96.97%，且三种物料的结果较均衡。每类仅 1 个原型时，煤矸石和矿渣准确率分别降至 83.10% 和 80.68%，说明单个典型外观不足以覆盖两类物料的现场变化。增加到每类 5 个原型后，煤矸石提高至 94.37%，但矿渣降至 80.68%，其主要错误集中为矿渣被识别为钢渣（17 张）。因此，当前数据和模型设置下，每类 3 个原型是准确率与物料间均衡性最好的选择；不能仅因煤矸石局部提高而选用 5 个原型。",
    )

    h2 = add_before(document, conclusion, "裁剪与视角变化实验", style="Heading 2")
    h2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_before(
        document,
        conclusion,
        "在增强 seed 42 数据集和每类 3 个原型的基础上，比较原始六视图 0.78 裁剪、仅使用全图、全图加中心视图及六视图 0.90 裁剪四种策略。该实验用于判断模型是否需要通过更多局部视角提取细节，或通过更大裁剪区域保留更多堆体整体形态。",
    )
    add_table(
        document,
        conclusion,
        ["编号", "视角设置", "裁剪比例", "原始准确率", "拒识率", "正确接受率", "煤矸石→钢渣"],
        [
            ["V0", "全图+中心+四角", "0.78", "96.97%", "21.21%", "76.62%", "5 张"],
            ["V1", "仅全图", "不裁剪", "91.77%", "19.48%", "75.76%", "12 张"],
            ["V2", "全图+中心", "0.78", "94.37%", "11.69%", "83.55%", "9 张"],
            ["V3", "全图+中心+四角", "0.90", "94.81%", "20.35%", "76.62%", "6 张"],
        ],
        "表6 裁剪与视角变化实验的独立测试结果",
    )
    add_figure(document, conclusion, view_figure, "图11 不同裁剪与视角策略的总体准确率及煤矸石误识为钢渣情况")
    add_before(
        document,
        conclusion,
        "原始六视图 0.78 裁剪策略（V0）仍是直接三分类性能最好的设置，总体原始准确率为 96.97%，并且煤矸石被误识为钢渣的数量最少。仅使用全图（V1）时准确率下降 5.20 个百分点，表明局部视角有助于模型提取物料纹理和颗粒细节。全图加中心视图（V2）的拒识率较低、正确接受率较高，但其原始准确率较 V0 低 2.60 个百分点，同时错误接受数量增加，因此不能据此判定模型分类更优。将裁剪比例增大到 0.90（V3）也没有带来提升。由此保留 V0 作为后续统一三分类模型的默认视角策略。",
    )
    add_before(
        document,
        conclusion,
        "阶段性结论：在固定独立测试集上，轻度训练图像增强带来可重复的总体提升；每类 3 个原型和六视图 0.78 裁剪在现有实验中取得最好的综合结果。剩余难点集中在煤矸石与钢渣的稳定混淆，后续应以这一错误方向为目标开展有针对性的样本分组和专门分支实验，而不宜继续仅靠叠加常规参数变化。",
    )
    document.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
