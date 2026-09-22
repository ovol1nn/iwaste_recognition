"""Normalize the legacy metric label in the 260922 report tables."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX_PATH = Path(__file__).resolve().parents[2] / "新峰测验" / "260922 识别模型实现说明.docx"


def format_header(cell) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run("自动正确\n输出比例")
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(8.5)
    run.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def main() -> None:
    document = Document(DOCX_PATH)
    updated = 0
    for table in document.tables:
        for cell in table.rows[0].cells:
            if "正确接受率" in cell.text:
                format_header(cell)
                updated += 1
    if updated != 2:
        raise ValueError(f"expected 2 legacy headers, updated {updated}")
    document.save(DOCX_PATH)
    print(f"updated_headers={updated}")


if __name__ == "__main__":
    main()
