"""Append a per-material baseline/augmentation comparison table to the model note."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DOCX_PATH = PROJECT_ROOT / "新峰测验" / "260908 识别模型实现说明.docx"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def format_run(run, size: float = 8.8, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def insert_paragraph_before(anchor, document: Document, text: str = "", align=None):
    paragraph = document.add_paragraph()
    if text:
        format_run(paragraph.add_run(text), size=10.5)
    if align is not None:
        paragraph.alignment = align
    anchor._p.addprevious(paragraph._p)
    return paragraph


def insert_material_table(document: Document, anchor) -> None:
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["物料（测试数）", "基线", "增强\nseed 42", "增强\nseed 2026", "增强\nseed 2027", "增强均值", "提升"]
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        format_run(run, bold=True, color="FFFFFF")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "1F4E78")

    rows = [
        ("煤矸石（n=71）", "90.14%", "92.96%", "91.55%", "91.55%", "92.02%", "+1.88 个百分点"),
        ("矿渣（n=88）", "95.45%", "98.86%", "96.59%", "96.59%", "97.35%", "+1.90 个百分点"),
        ("钢渣（n=72）", "91.67%", "98.61%", "95.83%", "94.44%", "96.29%", "+4.62 个百分点"),
    ]
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for cell, text in zip(cells, values, strict=True):
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            format_run(run, size=8.8, bold=row_index == 3)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 0:
                set_cell_shading(cell, "EAF2F8")
    table.autofit = True
    anchor._p.addprevious(table._tbl)


def main() -> None:
    document = Document(DOCX_PATH)
    if any("三种物料在基线与增强实验中的原始识别准确率" in paragraph.text for paragraph in document.paragraphs):
        raise ValueError("material comparison table already exists")
    anchor = next((paragraph for paragraph in document.paragraphs if paragraph.text.startswith("按物料分析")), None)
    if anchor is None:
        raise ValueError("cannot locate the per-material analysis paragraph")

    insert_paragraph_before(
        anchor,
        document,
        "进一步按物料拆分测试结果，可以区分增强策略对不同物料的实际作用。三次增强实验使用同一测试集，表中均为原始识别准确率。",
    )
    insert_material_table(document, anchor)
    caption = insert_paragraph_before(
        anchor,
        document,
        "表5 三种物料在基线与增强实验中的原始识别准确率（测试集）",
        WD_ALIGN_PARAGRAPH.CENTER,
    )
    format_run(caption.runs[0], size=10)
    document.save(DOCX_PATH)


if __name__ == "__main__":
    main()
